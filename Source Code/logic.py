import os
import re
import mimetypes
import shutil
import magic
import tempfile
import logging
import zipfile
import xml.etree.ElementTree as ET

from PIL import Image, ImageOps
from PIL.ExifTags import TAGS, GPSTAGS
import docx
import pptx
import openpyxl
import PyPDF2
from oletools.olevba import VBA_Parser
import vt

logger = logging.getLogger(__name__)


def get_file_size(filepath):
    try:
        return os.path.getsize(filepath)
    except Exception:
        return "Unknown"

def get_real_mime_and_ext(filepath):
    try:
        mime = magic.Magic(mime=True)
        real_mime = mime.from_file(filepath)
        real_ext = mimetypes.guess_extension(real_mime) or ""
        
        if real_mime in ["application/x-dosexec", 
                        "application/x-msdownload",
                        "application/x-executable",
                        "application/x-msdos-program",
                        "application/x-winexe"]:
            real_ext = ".exe"
        
        return real_mime, real_ext
    except Exception:
        return None, ""

# Extensions that are all valid for the same real file type. mimetypes only
# returns one of them, so without this a ".jpeg" photo was flagged as a
# mismatch because the detected extension was ".jpg".
EXTRA_EXTENSIONS = {
    "image/jpeg": {".jpg", ".jpeg", ".jpe", ".jfif"},
    "image/tiff": {".tif", ".tiff"},
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": {".docx", ".docm", ".dotx", ".dotm"},
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": {".xlsx", ".xlsm", ".xltx", ".xltm"},
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": {".pptx", ".pptm", ".potx", ".ppsx", ".ppsm"},
}

# Plain text has no fixed extension (.txt, .md, .csv, .log, ...), so a text
# file is never treated as a mismatch.
NO_MISMATCH_MIMES = {"text/plain", "text/csv"}

def is_extension_mismatch(ext, real_mime, real_ext):
    if not real_mime or not real_ext or real_ext == "Unknown":
        return False
    if real_mime in NO_MISMATCH_MIMES:
        return False
    valid = set(mimetypes.guess_all_extensions(real_mime))
    valid.add(real_ext)
    valid |= EXTRA_EXTENSIONS.get(real_mime, set())
    return ext.lower() not in valid

def sanitize_filename(filename):
    return re.sub(r'[^a-zA-Z0-9._-]', '_', filename)

def prefix_filename(filename, threat_level):
    prefix = "SAFE_" if threat_level == "Low" else "HIGH_RISK_"
    name, ext = os.path.splitext(filename)
    return prefix + name + ext

def avoid_duplicate(filepath):
    dirname, basename = os.path.split(filepath)
    name, ext = os.path.splitext(basename)
    counter = 1
    new_filepath = filepath
    while os.path.exists(new_filepath):
        new_name = f"{name}({counter}){ext}"
        new_filepath = os.path.join(dirname, new_name)
        counter += 1
    return new_filepath

def ensure_dir(filepath):
    dirname = os.path.dirname(filepath)
    if dirname and not os.path.exists(dirname):
        os.makedirs(dirname)


def to_float(value):
    # Pillow returns EXIF fractions as IFDRational; older files may give (num, den) tuples.
    if isinstance(value, tuple) and len(value) == 2:
        return value[0] / value[1]
    return float(value)

def convert_gps_to_degrees(value):
    try:
        d, m, s = (to_float(v) for v in value)
        return d + (m / 60.0) + (s / 3600.0)
    except Exception:
        return None

def extract_gps_data(exif_data):
    gps_metadata = {}
    
    try:
        gps_ifd = exif_data.get_ifd(0x8825)
        
        if gps_ifd:
            gps_data = {}
            
            for tag, value in gps_ifd.items():
                tag_name = GPSTAGS.get(tag, tag)
                gps_data[tag_name] = value
            
            if 'GPSLatitude' in gps_data and 'GPSLongitude' in gps_data:
                lat = convert_gps_to_degrees(gps_data['GPSLatitude'])
                lon = convert_gps_to_degrees(gps_data['GPSLongitude'])
                
                if lat is not None and lon is not None:
                    lat_ref = gps_data.get('GPSLatitudeRef', 'N')
                    lon_ref = gps_data.get('GPSLongitudeRef', 'E')

                    gps_metadata['GPS_Latitude'] = f"{lat:.6f}° {lat_ref}"
                    gps_metadata['GPS_Longitude'] = f"{lon:.6f}° {lon_ref}"

                    if lat_ref == 'S':
                        lat = -lat
                    if lon_ref == 'W':
                        lon = -lon
                    gps_metadata['GPS_Coordinates'] = f"{lat:.6f}, {lon:.6f}"

            # Each field is read separately so one bad value does not hide the rest.
            if 'GPSAltitude' in gps_data:
                try:
                    gps_metadata['GPS_Altitude'] = f"{to_float(gps_data['GPSAltitude']):.2f} meters"
                except Exception:
                    pass

            if 'GPSDateStamp' in gps_data:
                gps_metadata['GPS_Date'] = str(gps_data['GPSDateStamp'])

            if 'GPSTimeStamp' in gps_data:
                try:
                    h, m, s = (int(to_float(v)) for v in gps_data['GPSTimeStamp'])
                    gps_metadata['GPS_Time'] = f"{h:02d}:{m:02d}:{s:02d} UTC"
                except Exception:
                    pass

    except Exception:
        pass

    return gps_metadata


def parse_pdf_date(pdf_date_string):
    try:
        date_str = str(pdf_date_string).strip()
        if date_str.startswith('D:'):
            date_str = date_str[2:]
        
        year = date_str[0:4]
        month = date_str[4:6]
        day = date_str[6:8]
        hour = date_str[8:10] if len(date_str) >= 10 else '00'
        minute = date_str[10:12] if len(date_str) >= 12 else '00'
        second = date_str[12:14] if len(date_str) >= 14 else '00'
        
        readable = f"{year}-{month}-{day} {hour}:{minute}:{second}"
        
        if '+' in date_str or '-' in date_str:
            tz_match = re.search(r"([+-]\d{2})'(\d{2})'?", date_str)
            if tz_match:
                tz_hours = tz_match.group(1)
                tz_mins = tz_match.group(2)
                if tz_hours == '+00' and tz_mins == '00':
                    readable += " UTC"
                else:
                    readable += f" (UTC{tz_hours}:{tz_mins})"
        
        return readable
    
    except Exception:
        return str(pdf_date_string)

def clean_pdf_metadata_key(key):
    clean_key = str(key).lstrip('/')
    clean_key = re.sub(r'([a-z])([A-Z])', r'\1 \2', clean_key)
    
    key_mapping = {
        'ModDate': 'Modified Date',
        'CreationDate': 'Creation Date',
        'PTEX.Fullbanner': 'Generator'
    }
    
    return key_mapping.get(clean_key, clean_key)


EXIF_IFD = 0x8769
GPS_IFD = 0x8825
MAKER_NOTE = 0x927C

def add_exif_tags(metadata, tags):
    for tag_id, value in tags.items():
        if tag_id in (EXIF_IFD, GPS_IFD, MAKER_NOTE):
            continue
        if value is None or isinstance(value, bytes):
            continue
        value_str = str(value).strip().strip('\x00')
        if value_str and value_str.lower() != 'none':
            metadata[str(TAGS.get(tag_id, tag_id))] = value_str

def extract_image_metadata(filepath):
    metadata = {}
    try:
        with Image.open(filepath) as image:
            exifdata = image.getexif()

            # Main tags (camera make, model, artist, ...)
            add_exif_tags(metadata, exifdata)

            # Exif sub-section (date taken, camera serial number, lens, ...)
            try:
                add_exif_tags(metadata, exifdata.get_ifd(EXIF_IFD))
            except Exception:
                pass

            gps_data = extract_gps_data(exifdata)
            if gps_data:
                metadata.update(gps_data)

            # Text chunks such as PNG "Author" or "Comment"
            for key, value in image.info.items():
                if isinstance(value, str) and value.strip() and key not in metadata:
                    metadata[str(key)] = value.strip()

    except Exception:
        pass

    return metadata


# Office files (.docx, .pptx, .xlsx) keep their metadata in the same three
# XML parts inside the zip, so they are read and cleaned the same way.
CORE_PART = "docProps/core.xml"
APP_PART = "docProps/app.xml"
CUSTOM_PART = "docProps/custom.xml"

# Fields in app.xml that can identify a person or organisation.
APP_SENSITIVE_FIELDS = ["Company", "Manager", "HyperlinkBase", "Template"]

CORE_FIELD_NAMES = {
    "creator": "author",
    "lastModifiedBy": "last_modified_by",
    "description": "comments",
    "contentStatus": "content_status",
    "lastPrinted": "last_printed",
}

EMPTY_CORE_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<cp:coreProperties '
    'xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
    'xmlns:dc="http://purl.org/dc/elements/1.1/" '
    'xmlns:dcterms="http://purl.org/dc/terms/" '
    'xmlns:dcmitype="http://purl.org/dc/dcmitype/" '
    'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance"/>'
)

EMPTY_CUSTOM_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
    'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"/>'
)

def local_name(tag):
    return tag.rsplit('}', 1)[-1]

def extract_office_metadata(filepath):
    metadata = {}
    try:
        with zipfile.ZipFile(filepath) as z:
            names = set(z.namelist())

            if CORE_PART in names:
                for el in ET.fromstring(z.read(CORE_PART)):
                    value = (el.text or "").strip()
                    if value:
                        name = local_name(el.tag)
                        metadata[CORE_FIELD_NAMES.get(name, name)] = value.replace("T", " ").rstrip("Z") if name in ("created", "modified", "lastPrinted") else value

            if APP_PART in names:
                for el in ET.fromstring(z.read(APP_PART)):
                    name = local_name(el.tag)
                    value = (el.text or "").strip()
                    if name in APP_SENSITIVE_FIELDS and value:
                        metadata[name] = value

            if CUSTOM_PART in names:
                for prop in ET.fromstring(z.read(CUSTOM_PART)):
                    prop_name = prop.get("name")
                    value = "".join(prop.itertext()).strip()
                    if prop_name and value:
                        metadata[f"custom: {prop_name}"] = value
    except Exception:
        pass
    return metadata

def clean_office_file(src_path, dst_path):
    with zipfile.ZipFile(src_path) as zin, zipfile.ZipFile(dst_path, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == CORE_PART:
                data = EMPTY_CORE_XML.encode("utf-8")
            elif item.filename == CUSTOM_PART:
                data = EMPTY_CUSTOM_XML.encode("utf-8")
            elif item.filename == APP_PART:
                text = data.decode("utf-8")
                for field in APP_SENSITIVE_FIELDS:
                    text = re.sub(rf"<{field}>.*?</{field}>", f"<{field}/>", text, flags=re.S)
                data = text.encode("utf-8")
            zout.writestr(item, data)

def extract_pdf_metadata(filepath):
    metadata = {}
    try:
        with open(filepath, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            
            if reader.metadata:
                for key, value in reader.metadata.items():
                    if value is not None:
                        clean_key = clean_pdf_metadata_key(key)
                        value_str = str(value).strip()
                        
                        if value_str and value_str.lower() != 'none':
                            if 'date' in clean_key.lower() and value_str.startswith('D:'):
                                value_str = parse_pdf_date(value_str)
                            
                            metadata[clean_key] = value_str
    
    except Exception:
        pass
    
    return metadata


OFFICE_MIMES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

def check_office_type(filepath):
    try:
        doc = docx.Document(filepath)
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".docx"
    except Exception:
        pass
        
    try:
        prs = pptx.Presentation(filepath)
        return "application/vnd.openxmlformats-officedocument.presentationml.presentation", ".pptx"
    except Exception:
        pass
        
    try:
        wb = openpyxl.load_workbook(filepath, read_only=True)
        wb.close()
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", ".xlsx"
    except Exception:
        pass
        
    return None, None


def detect_macros(filepath):
    vbaparser = None
    try:
        vbaparser = VBA_Parser(filepath)
        if vbaparser.detect_vba_macros():
            return True, "VBA Macros found!"
        else:
            return False, "No VBA macros found."
    except Exception:
        return False, "Macro detection error."
    finally:
        if vbaparser:
            vbaparser.close()


def analyze_file(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    file_size = get_file_size(filepath)
    mime_type, _ = mimetypes.guess_type(filepath)
    
    metadata = {}
    
    office_mime, office_ext = check_office_type(filepath)
    
    if office_mime:
        real_mime, real_ext = office_mime, office_ext
    else:
        real_mime, real_ext = get_real_mime_and_ext(filepath)
    

    if real_mime:
        if real_mime.startswith("image/"):
            metadata = extract_image_metadata(filepath)
        elif real_mime == "application/pdf":
            metadata = extract_pdf_metadata(filepath)
        elif real_mime in OFFICE_MIMES:
            metadata = extract_office_metadata(filepath)
    
    meta_str = "\n".join(f"{k}: {v}" for k, v in metadata.items()) if metadata else "No metadata found."
    
    office_mimes = [
        "application/vnd.ms-excel",
        "application/vnd.ms-powerpoint",
        "application/msword",
        "application/vnd.openxmlformats-officedocument"
    ]
    macro_found, macro_str = (False, "No macros or scripts detected.")
    
    if real_mime and any(real_mime.startswith(m) for m in office_mimes):
        macro_found, macro_str = detect_macros(filepath)
    
    mismatch = is_extension_mismatch(ext, real_mime, real_ext)
    threat_level = "High" if macro_found or mismatch else "Low"
    
    summary = []
    if mismatch:
        summary.append(f"File type mismatch: Extension is {ext} but file is {real_ext}.")
    summary.append(f"Threat Level: {threat_level.upper()}")
    
    return {
        "file_size": file_size,
        "current_extension": ext,
        "mime_type": mime_type or "Unknown",
        "real_mime_type": real_mime or "Unknown",
        "real_extension": real_ext or "Unknown",
        "metadata": meta_str,
        "macro_detection": macro_str,
        "macro_found": macro_found,
        "extension_mismatch": mismatch,
        "summary": "\n".join(summary),
        "threat_level": threat_level
    }


def write_clean_file(src_path, dst_path, analysis=None):
    if analysis is None:
        analysis = analyze_file(src_path)
    
    real_mime = analysis["real_mime_type"]
    
    handled = False
    
    try:
        if real_mime in ("image/jpeg", "image/png"):
            with Image.open(src_path) as image:
                image.load()
                # Apply the EXIF rotation to the pixels first, otherwise the
                # photo shows sideways once the rotation tag is removed.
                clean = ImageOps.exif_transpose(image)
                save_args = {}
                # The colour profile and transparency are not personal data
                # and are needed for the picture to look the same.
                if image.info.get("icc_profile"):
                    save_args["icc_profile"] = image.info["icc_profile"]
                if real_mime == "image/png" and "transparency" in image.info:
                    save_args["transparency"] = image.info["transparency"]
                clean.info = {}
                if real_mime == "image/jpeg":
                    clean.save(dst_path, format="JPEG", quality=95, **save_args)
                else:
                    clean.save(dst_path, format="PNG", **save_args)
            handled = True

        elif real_mime == "application/pdf":
            with open(src_path, 'rb') as f_in:
                reader = PyPDF2.PdfReader(f_in)
                writer = PyPDF2.PdfWriter()
                for page in reader.pages:
                    writer.add_page(page)
                # PdfWriter adds "Producer: PyPDF2" by default.
                writer.add_metadata({"/Producer": ""})
                with open(dst_path, 'wb') as f_out:
                    writer.write(f_out)
            handled = True

        elif real_mime in OFFICE_MIMES:
            clean_office_file(src_path, dst_path)
            handled = True

        if not handled:
            shutil.copy2(src_path, dst_path)
            return True, "File type not supported for metadata removal. Copied as-is."

        return True, "Metadata removed from file."

    except Exception as e:
        logger.error(f"Failed to clean file: {e}")
        return False, f"Failed to create clean file: {e}"

def remove_metadata_in_place(filepath):
    temp_fd, temp_path = tempfile.mkstemp()
    os.close(temp_fd)

    try:
        success, msg = write_clean_file(filepath, temp_path)
        if not success:
            return False, msg
        shutil.move(temp_path, filepath)
        return True, msg
    except Exception as e:
        return False, str(e)
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


def export_safe_file(src, dst, threat_level, mime_type):
    analysis_result = analyze_file(src)
    real_ext = analysis_result.get("real_extension")
    if not real_ext or real_ext == "Unknown":
        real_ext = os.path.splitext(src)[1]

    dst_dir = os.path.dirname(dst)
    dst_name_no_ext = os.path.splitext(os.path.basename(dst))[0]
    final_dst = os.path.join(dst_dir, dst_name_no_ext + real_ext)
    final_path = avoid_duplicate(final_dst)
    ensure_dir(final_path)
    
    try:
        success, msg = write_clean_file(src, final_path, analysis_result)
        if success:
            export_msg = f"File exported as: {final_path}\n({msg})"
            return True, export_msg
        else:
            return False, f"Export failed during cleaning: {msg}"
    except Exception as e:
        return False, str(e)


def fix_extension_in_place(filepath):
    try:
        current_ext = os.path.splitext(filepath)[1]
        
        analysis_result = analyze_file(filepath)
        real_ext = analysis_result.get("real_extension")

        if not analysis_result.get("extension_mismatch") or real_ext == current_ext:
            return (False, "File extension is already correct or unknown.", filepath)
        
        dirname = os.path.dirname(filepath)
        basename_no_ext = os.path.splitext(os.path.basename(filepath))[0]
        new_filepath = os.path.join(dirname, basename_no_ext + real_ext)
        
        final_path = avoid_duplicate(new_filepath)
        os.rename(filepath, final_path)
        return (True, f"File renamed to {os.path.basename(final_path)}", final_path)
    except Exception as e:
        return (False, str(e), filepath)


def scan_with_virustotal(filepath, api_key):
    client = None
    try:
        client = vt.Client(api_key)
        
        with open(filepath, 'rb') as f:
            analysis = client.scan_file(f, wait_for_completion=True)
        
        if analysis.status != "completed":
            return {
                "status": "Error", 
                "message": f"Scan failed. Status: {analysis.status}", 
                "details": ""
            }

        stats = analysis.stats
        positives = stats.get('malicious', 0) + stats.get('suspicious', 0)
        total_scans = stats.get('total', 
                                stats.get('harmless', 0) + 
                                stats.get('malicious', 0) + 
                                stats.get('suspicious', 0) + 
                                stats.get('undetected', 0))

        if total_scans == 0:
            return {
                "status": "Error", 
                "message": "Scan complete, but no engine results.", 
                "details": ""
            }

        status = "Safe"
        details_list = []
        if positives > 0:
            if positives <= 5:
                status = "Suspicious"
            else:
                status = "HighRisk"
            
            for result in analysis.results.values():
                if result['category'] == 'malicious' or result['category'] == 'suspicious':
                    if result['result']:
                        details_list.append(f"- {result['engine_name']}: {result['result']}")
        
        message = f"{positives} / {total_scans} engines detected this file."
        
        details_str = ""
        if details_list:
            details_str = "\n".join(details_list[:5])
            if len(details_list) > 5:
                details_str += f"\n...and {len(details_list) - 5} more."
        elif positives > 0:
            details_str = "Detections found, but no names reported."

        return {
            "status": status,
            "message": message,
            "details": details_str
        }
    
    except vt.error.APIError as e:
        return {"status": "Error", "message": f"API Error: {e.message}", "details": ""}
    except Exception as e:
        return {"status": "Error", "message": f"Scan Error: {str(e)}", "details": ""}
    finally:
        if client:
            client.close()